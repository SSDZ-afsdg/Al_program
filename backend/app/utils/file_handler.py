"""
file_handler.py —— 合同文件读取与文本提取工具

支持两种格式：
- .docx：使用 python-docx 按段落与表格提取文字；
- .pdf ：使用 PyPDF2 逐页提取文字。
"""

from pathlib import Path
from typing import Optional

from app.core.exceptions import BizException


def extract_text_from_docx(file_path: str | Path) -> str:
    """
    解析 .docx Word 文档并提取纯文本。

    :param file_path: docx 文件路径
    :return: 提取出的纯文本（段落之间以换行连接）
    """
    # 延迟导入：只有真正处理 docx 时才加载库，减少其他场景的启动开销
    from docx import Document as DocxDocument

    # 打开 Word 文档
    doc = DocxDocument(str(file_path))

    text_parts = []

    # 1. 提取正文段落（跳过空段落）
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped:
            text_parts.append(stripped)

    # 2. 提取表格中的文字（合同中的关键信息常以表格呈现）
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # 同一行单元格用空格拼接
                text_parts.append(" ".join(row_cells))

    return "\n".join(text_parts)


def extract_text_from_pdf(file_path: str | Path) -> str:
    """
    解析 PDF 文档并提取纯文本。

    :param file_path: pdf 文件路径
    :return: 提取出的纯文本（按页拼接）
    """
    # 延迟导入 PyPDF2
    from PyPDF2 import PdfReader

    # 以二进制只读方式打开 PDF
    reader = PdfReader(str(file_path))

    text_parts = []
    # 逐页提取文本
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text and page_text.strip():
            text_parts.append(page_text.strip())

    return "\n".join(text_parts)


def extract_text(file_path: str | Path) -> str:
    """
    根据文件扩展名自动选择解析器，提取合同文本。

    :param file_path: 文件路径
    :return: 提取出的纯文本
    :raises BizException: 文件格式不支持或文件为空时抛出 400 异常
    """
    path = Path(file_path)
    # 统一转小写取扩展名（不含点），如 ".DOCX" -> "docx"
    suffix: Optional[str] = path.suffix.lower().lstrip(".") if path.suffix else None

    # 按扩展名分派到对应的解析函数
    if suffix == "docx":
        text = extract_text_from_docx(path)
    elif suffix == "pdf":
        text = extract_text_from_pdf(path)
    else:
        raise BizException(code=400, message=f"不支持的文件格式：.{suffix}，仅支持 .docx / .pdf")

    # 去除首尾空白后校验：防止上传无文字的空白文档或扫描件 PDF
    text = text.strip()
    if not text:
        raise BizException(
            code=400,
            message="未能从文件中提取到文本，请确认文件非空白文档（扫描件 PDF 暂不支持 OCR 识别）",
        )

    return text
